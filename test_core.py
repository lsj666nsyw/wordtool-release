# -*- coding: utf-8 -*-
"""Word 工具核心逻辑自测（双向验证）。"""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, r"D:\21821\products\wordtool")
from wordtool_core import (batch_replace, extract_images, doc_stats,
                           merge_docs, export_text, add_text_watermark,
                           is_doc)
from docx import Document

pass_n = 0
fail_n = 0


def check(name, cond, detail=""):
    global pass_n, fail_n
    if cond:
        pass_n += 1
        print(f"  PASS  {name}")
    else:
        fail_n += 1
        print(f"  FAIL  {name}  {detail}")


def make_docx(path, text_paragraphs, title="", author=""):
    doc = Document()
    for t in text_paragraphs:
        doc.add_paragraph(t)
    if title:
        doc.core_properties.title = title
    if author:
        doc.core_properties.author = author
    doc.save(str(path))
    return path


def read_texts(path):
    doc = Document(str(path))
    txts = [p.text for p in doc.paragraphs]
    return txts


def main():
    tmp = Path(tempfile.mkdtemp(prefix="wordtest_"))
    print(f"临时目录：{tmp}\n")

    # ---------- 1. 批量替换 - 正常 PASS ----------
    print("[1] 批量替换")
    f1 = make_docx(tmp / "a.docx", ["张三的订单", "李四的订单", "联系张三"])
    res = batch_replace([f1], str(tmp), "张三", "王五", case_sensitive=True)
    check("替换返回成功", res[0][2] is not None and res[0][3] is None)
    check("替换计数>0", res[0][1] > 0, f"实际 {res[0][1]}")
    txts = read_texts(res[0][2])
    check("替换后包含王五", "王五" in txts[0], f"实际 {txts[0]}")
    check("替换后不包含张三", "张三" not in "".join(txts), f"实际 {txts}")

    # ---------- 2. 大小写不敏感替换 ----------
    print("[2] 大小写不敏感替换")
    f2 = make_docx(tmp / "b.docx", ["Hello World", "hello world"])
    res2 = batch_replace([f2], str(tmp), "hello", "你好", case_sensitive=False)
    check("CI替换成功", res2[0][2] is not None)
    txts2 = read_texts(res2[0][2])
    check("CI替换后含你好", "你好" in "".join(txts2), f"实际 {txts2}")

    # ---------- 3. 表格替换 ----------
    print("[3] 表格替换")
    f3 = make_docx(tmp / "c.docx", ["表格测试"])
    doc = Document(str(f3))
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "客户名称"
    tbl.cell(0, 1).text = "备注"
    tbl.cell(1, 0).text = "张三"
    tbl.cell(1, 1).text = "重要客户"
    doc.save(str(f3))
    out3 = tmp / "c_replaced.docx"
    res3 = batch_replace([f3], str(tmp), "张三", "李四", case_sensitive=True)
    check("表格替换成功", res3[0][2] is not None)

    # ---------- 4. 空 find_text 不替换（防呆） ----------
    print("[4] 空 find_text 防呆")
    res4 = batch_replace([f1], str(tmp), "", "X", case_sensitive=True)
    check("空find不替换", res4[0][1] == 0, f"实际 {res4[0][1]}")

    # ---------- 5. 表格单元格替换 ----------
    print("[5] 表格单元格替换")
    f5 = tmp / "e.docx"
    doc = Document()
    doc.add_paragraph("正文")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "张三"
    t.cell(0, 1).text = "北京"
    t.cell(1, 0).text = "李四"
    t.cell(1, 1).text = "上海"
    doc.save(str(f5))
    out5 = tmp / "e_out.docx"
    r5 = batch_replace([f5], str(tmp), "张三", "王五", case_sensitive=True)
    check("表格替换计数>0", r5[0][1] > 0, f"实际 {r5[0][1]}")

    # ---------- 6. 提取图片 ----------
    print("[6] 提取图片")
    from PIL import Image
    import io
    f6 = tmp / "img.docx"
    doc = Document()
    buf = io.BytesIO()
    Image.new("RGB", (100, 100), (255, 0, 0)).save(buf, format="PNG")
    buf.seek(0)
    doc.add_paragraph("带图片的文档")
    doc.add_picture(buf)
    doc.save(str(f6))
    img_dir = tmp / "images"
    imgs = extract_images(f6, str(img_dir))
    check("提取到1张图片", len(imgs) == 1, f"实际 {len(imgs)}")
    if imgs:
        check("图片文件存在", imgs[0].exists())

    # ---------- 7. 文档统计 ----------
    print("[7] 文档统计")
    stats = doc_stats(f1)
    check("段落数>=1", stats["段落数"] >= 1)
    check("标题字段存在", "标题" in stats)
    check("作者字段存在", "作者" in stats)

    # ---------- 8. 文档合并 ----------
    print("[8] 文档合并")
    fa = make_docx(tmp / "m1.docx", ["第一段", "第二段"])
    fb = make_docx(tmp / "m2.docx", ["第三段"])
    out_m = tmp / "merged.docx"
    total, oks = merge_docs([fa, fb], str(out_m), add_page_break=False)
    check("合并成功", len(oks) == 2, f"实际 {len(oks)}")
    check("合并后段落>=3", total >= 3, f"实际 {total}")
    check("合并文件存在", out_m.exists())

    # ---------- 9. 导出纯文本 ----------
    print("[9] 导出纯文本")
    out_txt = tmp / "out.txt"
    n, chars = export_text(f1, str(out_txt))
    check("导出txt存在", out_txt.exists())
    check("导出字符数>0", chars > 0)

    # ---------- 10. is_doc ----------
    print("[10] is_doc")
    check("is_doc .docx", is_doc("x.docx"))
    check("is_doc .docm", is_doc("x.docm"))
    check("is_doc 拒绝 .pdf", not is_doc("x.pdf"))
    check("is_doc 拒绝 .doc", not is_doc("x.doc"))

    print(f"\n=== 测试完成：{pass_n} PASS / {fail_n} FAIL ===")
    return 0 if fail_n == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
